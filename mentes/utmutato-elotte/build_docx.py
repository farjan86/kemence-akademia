# -*- coding: utf-8 -*-
"""Kemence Akadémia - Felhasználói útmutató -> .docx"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# alap betűméret
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    pPr.append(shd)

def segs(p, segments):
    for seg in segments:
        text, bold = (seg if isinstance(seg, tuple) else (seg, False))
        r = p.add_run(text); r.bold = bold

def para(segments, style=None):
    p = doc.add_paragraph(style=style); segs(p, segments); return p

def callout(segments, fill='FCEEE4'):
    p = doc.add_paragraph(); segs(p, segments); shade(p, fill)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    return p

def screenshot(caption):
    p = doc.add_paragraph(); shade(p, 'F3E9D8')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('📷  Képernyőkép: ' + caption); r.italic = True
    r.font.color.rgb = RGBColor(0x8A, 0x74, 0x58)
    return p

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.paragraphs[0].add_run(h).bold = True
        shade(c.paragraphs[0], 'F6EFE6')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''; cells[i].paragraphs[0].add_run(val)
    doc.add_paragraph()
    return t

def h1(text): doc.add_heading(text, level=1)
def h2(text): doc.add_heading(text, level=2)

# ===================== CÍM =====================
title = doc.add_heading('', level=0)
title.add_run('Kemence Akadémia — Felhasználói útmutató')
doc.add_paragraph('A foglalási rendszer használata: a vendég oldala és az admin felület, lépésről lépésre.')
p = doc.add_paragraph(); r = p.add_run('A képernyőképek helye jelölve (📷) — ezeket később illesztitek be.'); r.italic = True

callout([('Miből áll a rendszer? ', True),
         'Publikus oldal — ezt látják a vendégek (böngésznek, foglalnak). Admin felület (/admin/) — ezt kezelitek ti (foglalások, naptár, programok, beállítások). A leveleket a rendszer automatikusan vagy egy gombnyomásra küldi.'], 'EEF7F0')

# ===================== 1 =====================
h1('1. Hogyan foglal a vendég?')
para(['A vendég a publikus oldalon böngészi a programokat, és pár kattintással foglal. Nem kell regisztrálnia.'])
h2('1.1 A főoldal')
para(['A vendég a főoldalon látja a bemutatkozást, a workshop-típusokat, és a Programok résznél a meghirdetett, foglalható alkalmakat (cím, időpont, ár, szabad helyek). Minden programnál egy rövid leírás látszik; a „Részletek" gombra kattintva a hosszabb, részletes leírás is megjelenik (ha van megadva).'])
screenshot('a publikus főoldal a programlistával')
h2('1.2 A foglalás lépései')
para([('1. ', True), 'A vendég kiválaszt egy programot, és a Foglalás gombra kattint.'], style='List Number')
para([('2. ', True), 'Kitölti az űrlapot: név, e-mail, telefon, létszám (hány fő), megjegyzés (nem kötelező).'], style='List Number')
para([('3. ', True), 'Elküldi. Ha van elég szabad hely, a rendszer rögzíti a foglalást.'], style='List Number')
screenshot('a foglalási űrlap (felugró ablak)')
callout([('Túlfoglalás elleni védelem: ', True), 'a rendszer nem enged több foglalást, mint ahány szabad hely van, és múltbéli programra sem enged a vendég foglalni.'])
h2('1.3 Mi történik közvetlenül a foglalás után?')
para(['A foglalás bekerül a rendszerbe „Jóváhagyásra vár" státusszal.'], style='List Bullet')
para(['A vendég azonnal kap egy visszaigazoló levelet („megkaptuk a foglalásod").'], style='List Bullet')
para(['Ti értesítőt kaptok a csapat e-mail címére, hogy új foglalás érkezett.'], style='List Bullet')

# ===================== 2 =====================
h1('2. Egy foglalás életútja és a státuszok')
para(['Minden foglalásnak van egy státusza, ami megmutatja, hol tart. Négy státusz van:'])
table(['Státusz', 'Mit jelent', 'Foglal helyet?'], [
    ['Jóváhagyásra vár', 'Új foglalás, még nem döntöttetek róla.', 'Igen (fenntartja a helyet)'],
    ['Jóváhagyott', 'Elfogadtátok, a vendég helye biztos.', 'Igen'],
    ['Elutasított', 'Nem fogadtátok el (pl. betelt).', 'Nem (felszabadítja)'],
    ['Lemondott', 'A vendég lemondta, ti rögzítitek.', 'Nem (felszabadítja)'],
])
h2('2.1 Milyen lépések lehetségesek?')
para(['„Jóváhagyásra vár" állapotból: Jóváhagyás, Elutasítás, vagy Vendég lemondta.'], style='List Bullet')
para(['„Jóváhagyott" állapotból: Vendég lemondta, vagy Jóváhagyás visszavonása (vissza „vár" állapotba).'], style='List Bullet')
para(['„Elutasított" / „Lemondott": lezárt állapot, innen nincs további lépés.'], style='List Bullet')
h2('2.2 Szabad helyek')
para(['A szabad helyek száma = max. létszám − (a „jóváhagyásra vár" + „jóváhagyott" foglalások létszáma). Ha egy foglalást elutasítasz vagy lemondasz, a helye azonnal felszabadul, és újra foglalható.'])

# ===================== 3 =====================
h1('3. Milyen e-mailek mennek ki?')
para(['A rendszer hétféle levelet ismer. Egy részük automatikus, más részüket ti küldtök egy gombnyomással (a szöveget küldés előtt megnézhetitek és szerkeszthetitek).'])
table(['Levél', 'Kinek', 'Mikor', 'Hogyan'], [
    ['Visszaigazolás', 'a vendégnek', 'foglaláskor', 'automatikus'],
    ['Csapat-értesítő', 'nektek (csapat cím)', 'foglaláskor', 'automatikus'],
    ['Jóváhagyás', 'a vendégnek', 'amikor jóváhagyjátok', 'gombbal (✉ Levél)'],
    ['Elutasítás', 'a vendégnek', 'amikor elutasítjátok', 'gombbal (✉ Levél)'],
    ['Lemondás', 'a vendégnek', 'lemondás rögzítésekor', 'gombbal (✉ Levél)'],
    ['Emlékeztető', 'a vendégnek', 'a program előtti napon', 'automatikus (napi)'],
    ['Program elmarad', 'az érintett foglalóknak', 'ha egy programot elmarasztotok', 'egy lépésben, tömegesen'],
])
callout([('A „✉ Levél" gomb: ', True), 'a foglalás sorában megjelenik egy levél-gomb. Rákattintva egy ablak nyílik, ahol látod és szerkesztheted a levél szövegét, mielőtt elküldöd.'])
callout([('Válaszcím: ', True), 'a vendégnek küldött levelekre a válasz a Csapat e-mail címre érkezik (Beállítások), tehát a vendég egyszerűen „válasz"-t nyomhat, és nektek ír.'], 'EEF7F0')
screenshot('a „✉ Levél" ablak a szerkeszthető szöveggel')

# ===================== 4 =====================
h1('4. Admin belépés')
para([('1. ', True), 'Nyisd meg az oldal /admin/ címét (pl. kemence-akademia.xyz/admin/).'], style='List Number')
para([('2. ', True), 'Add meg az e-mail címet és jelszót, amit a rendszerhez kaptál.'], style='List Number')
para(['Önregisztráció nincs, csak a meghívott admin fiók léphet be. A felület négy fülből áll: Foglalások, Naptár, Programok, Beállítások.'])
screenshot('a belépő képernyő és a négy fül')

# ===================== 5 =====================
h1('5. Foglalások kezelése')
para(['A Foglalások fülön látod és intézed a beérkezett foglalásokat.'])
h2('5.1 A négy státusz-alfül')
para(['Fent négy alfül van, státuszonként (mindegyiknél zárójelben a darabszám): Jóváhagyásra vár, Jóváhagyott, Elutasított, Lemondott. Kattints egyre, és az adott állapotú foglalások listáját látod.'])
h2('5.2 Szűrés')
para(['A lista szűkíthető: program szerint, vendég neve szerint, és időállapot szerint (jövőbeli / ma / múltbéli).'])
h2('5.3 Egy foglalás sora')
para(['Minden sorban látod: azonosító, program + időpont, időállapot, vendég (név, telefon, e-mail), létszám, megjegyzés, státusz, és a műveletek. A műveletek a státusztól függenek (lásd 2.1).'])
screenshot('a Foglalások lista egy sorral és a műveletekkel')
h2('5.4 Mit tehetsz egy foglalással?')
para([('Jóváhagyás', True), ' — elfogadod; a hely biztossá válik. Utána a ✉ Levéllel küldhetsz jóváhagyó értesítőt.'], style='List Bullet')
para([('Elutasítás', True), ' — nem fogadod el; a hely felszabadul. A ✉ Levéllel küldhetsz elutasítót.'], style='List Bullet')
para([('Vendég lemondta', True), ' — ha a vendég jelezte, itt rögzíted; a hely felszabadul.'], style='List Bullet')
para([('Szerkesztés', True), ' — javíthatod a foglalás adatait (név, e-mail, telefon, létszám, megjegyzés, program).'], style='List Bullet')
para([('✉ Levél', True), ' — a státuszhoz illő levelet küldöd a vendégnek (szerkeszthető szöveggel).'], style='List Bullet')
h2('5.5 Excel-export')
para(['A szűrősávban a 📊 Excel gombbal letöltöd a jelenleg látott listát (aktív alfül + szűrők) Excel-fájlba: azonosító, program, időpont, vendég elérhetőségek, létszám, státusz, megjegyzés.'])

# ===================== 6 =====================
h1('6. Naptár nézet és jelentés')
para(['A Naptár fülön egy éves áttekintést kapsz a programokról, és egy programra kattintva látod a résztvevőket.'])
h2('6.1 Az éves nézet')
para(['Léptethetsz az évek között, és látod, melyik hónapban milyen programok vannak. Egy programra kattintva lenyílnak a részletei.'])
screenshot('a naptár éves nézete a programokkal')
h2('6.2 Egy program résztvevői')
para(['A kiválasztott programnál látod a foglalt/szabad helyeket, a jóváhagyott és váró létszámot, és egy táblázatot az élő foglalókról (jóváhagyott + jóváhagyásra váró): azonosító, vendég neve, telefon, e-mail, létszám, státusz.'])
h2('6.3 Jelentés (PDF)')
para(['A résztvevő-tábla fölött a 📄 Jelentés (PDF) gombbal letöltesz egy nyomtatható PDF-et a programról: a program adatai (a részletes leírás nélkül) + az élő foglalók táblája elérhetőséggel és összesítéssel. Kiválóan használható a helyszínen (ki jön, hány fővel, hogyan érhető el).'])
screenshot('a program résztvevő-táblája + a Jelentés gomb, és egy kész PDF')

# ===================== 7 =====================
h1('7. Programok kezelése')
para(['A Programok fülön veszed fel és kezeled a workshopokat. Két alfül: Aktuális és Archivált programok.'])
h2('7.1 Program-státuszok')
table(['Státusz', 'Mit jelent'], [
    ['Aktív', 'Élő, foglalható program. Kötelező hozzá ár + időpont + létszám. Megjelenik a főoldalon.'],
    ['Hamarosan', 'Előzetes/csalogató. Megjelenik, de még nem foglalható (az ár/időpont lehet üres).'],
    ['Elmaradt', 'A program elmarad (nem lesz megtartva). A felületen „elmarad" jelöléssel látszik.'],
])
callout([('Az „archiválás" külön dolog! ', True), 'Az archivált jelző FÜGGETLEN a státusztól, és azt jelenti: látszik-e a főoldalon. Egy program lehet egyszerre „aktív jellegű" ÉS archivált (levettük a főoldalról, de az adatok/foglalások megmaradnak).'])
h2('7.2 Új program felvitele')
para([('1. ', True), 'Programok fül → Új program gomb.'], style='List Number')
para([('2. ', True), 'Töltsd ki: cím, rövid leírás (KÖTELEZŐ — ez jelenik meg a kártyán), részletes leírás (OPCIONÁLIS — a „Részletek" ablakban jön elő; ha üresen hagyod, nincs „Részletek" gomb), ár (és opcionálisan kedvezményes ár), időpont, várható időtartam, max. létszám, és állítsd be a státuszt (aktív / hamarosan).'], style='List Number')
para([('3. ', True), 'Ha van kép, tölts fel fotót a programhoz (ez lesz a foglalós kártya képe).'], style='List Number')
para([('4. ', True), 'Mentés. Ha „aktív", azonnal megjelenik a főoldalon és foglalható.'], style='List Number')
screenshot('az új program űrlapja')
h2('7.3 Mit tehetsz egy programmal?')
para([('Szerkesztés', True), ' — bármelyik adat módosítható (ár, időpont, létszám, leírás, kép, státusz).'], style='List Bullet')
para([('Archiválás', True), ' — leveszi a főoldalról, de az adatok és a foglalások megmaradnak (az „Archivált" alfülre kerül). Bármikor visszaállítható.'], style='List Bullet')
para([('„Elmarad a program"', True), ' — a státuszt elmaradt-ra állítja, és egy lépésben értesítő levelet küld az érintett (élő) foglalóknak.'], style='List Bullet')
para([('Törlés', True), ' — végleg törli a programot. (Ha van hozzá foglalás, inkább archiválj.)'], style='List Bullet')
h2('7.4 Mit jelent az archiválás a gyakorlatban?')
para(['Egy megtartott, lezajlott programot érdemes archiválni: eltűnik a főoldalról (a vendégek már nem látják/foglalják), de a foglalások és a statisztika megmarad, és visszanézhető. Nem törlöd el az adatokat, csak „elteszed".'])

# ===================== 8 =====================
h1('8. Beállítások')
para(['A Beállítások fülön szabod testre a rendszert.'])
para([('Csapat e-mail cím', True), ' — ide érkeznek a csapat-értesítők (új foglalás), ÉS ez a vendég-levelek válaszcíme (Reply-To). Ide írjátok a valódi, olvasott postaládát.'], style='List Bullet')
para([('Foglalási infosáv', True), ' — egy rövid szöveg, ami a foglalási űrlapon jelenik meg (pl. tudnivalók, feltételek).'], style='List Bullet')
para([('Azonosító-formátum', True), ' — a foglalások azonosítójának előtagja és kezdő száma (pl. F- + 100 → F-100).'], style='List Bullet')
para([('E-mail sablonok', True), ' — a hétféle levél tárgya és szövege szerkeszthető. Behelyettesíthető mezők használhatók (pl. a vendég neve, a program, az időpont), amiket a rendszer kitölt küldéskor.'], style='List Bullet')
screenshot('a Beállítások fül (csapat cím, infosáv, azonosító, sablonok)')

# ===================== 9 =====================
h1('9. Főoldali képek cseréje')
para(['A program-fotókat az adminból cseréled (7.2 / 7.3). A főoldal díszítő képei (hero-diavetítés, csapatkép, logó, workshop-bemutató képek) viszont fájlok, ezeket a webtárhelyen cseréled.'])
para([('1. ', True), 'Nyisd meg a web/kepek/ mappát.'], style='List Number')
para([('2. ', True), 'Cseréld az új képet UGYANAZON A NÉVEN (fontos, különben eltűnik): hero-1.jpg, hero-2.jpg, hero-3.jpg (diavetítés); csapat.jpg (Rólunk); logo-mark.png (logó); a *_workshop.jpg és leanybucsu.jpg (workshop-képek).'], style='List Number')
para([('3. ', True), 'Töltsd fel a frissített mappát a tárhelyre (új deploy).'], style='List Number')
callout([('Méret: ', True), 'ne tölts fel több MB-os / 4K képet, mert lassítja az oldalt. Ajánlott: hero ~2000px széles, a többi ~1400-1800px, JPG ~200-400 KB.'], 'FDF7E3')

callout([('Röviden a napi rutin: ', True), 'új foglalás jön → e-mailben értesültök → az admin a Foglalások fülön jóváhagyja (és a ✉ Levéllel értesíti a vendéget) → a program előtti nap a rendszer emlékeztetőt küld → a helyszínre a Naptár → Jelentés (PDF) viszi a résztvevőlistát.'], 'EEF7F0')

doc.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'felhasznaloi-utmutato.docx'))
print('KESZ: felhasznaloi-utmutato.docx')
