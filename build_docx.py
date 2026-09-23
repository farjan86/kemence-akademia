# -*- coding: utf-8 -*-
"""Kemence Akadémia — Admin kézikönyv.

EGY forrásból (a lenti TARTALOM) készül a két változat, így mindig egyformák:
    felhasznaloi-utmutato.html   és   felhasznaloi-utmutato.docx

Futtatás a projekt gyökeréből:   python build_docx.py
(Kell hozzá:  python -m pip install python-docx)

A szövegben:  **félkövér**   és   `kód`   jelölés használható.
"""
import os, re, html
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GYOKER = os.path.dirname(os.path.abspath(__file__))
CIM = "Kemence Akadémia — Admin kézikönyv"
ALCIM = "Három rész, sok példa: mit hol találsz, és mit csinálj egy-egy helyzetben."
FRISSITVE = "Frissítve: 2026-09-21"

# ---------------------------------------------------------------------
#  TARTALOM — blokkok:
#    ("h1", cím)  fő rész         ("h2", cím)  helyzet
#    ("p", szöveg)                ("lepesek", [..])  számozott lépések
#    ("lista", [..])              ("tipp", szöveg)   💡 miért / jó tudni
#    ("figyelem", szöveg)         ("tabla", [fejléc], [[sor], ..])
#    ("pelda", cím, [sorok])      📋 végigvitt gyakorlati példa
#    ("kep", leírás)  képernyőkép helye
# ---------------------------------------------------------------------
TARTALOM = [
("bevezeto", [
    "**Nyilvános oldal** — a vendégek itt foglalnak egy meghirdetett időpontra, vagy kérnek egyedi ajánlatot.",
    "**Admin felület** (`/admin/`) — itt intézed mindezt. Ez az útmutató erről szól.",
]),
("p", "Az útmutató **három részből** áll: az elsőben eligazodsz a felületen és beállítod a rendszert, a másodikban a meghirdetett programokat és a rájuk érkező foglalásokat kezeled, a harmadikban az egyedi ajánlatkéréseket."),

("tabla", ["Ezt szeretném…", "Itt találod"], [
    ["Jóváhagyok egy foglalást", "**Foglalások** fül → „Jóváhagyásra vár” (2. rész)"],
    ["Új alkalmat hirdetek meg", "**Programok** fül → „Szerkesztés” → „+ Időpont hozzáadása” (2. rész)"],
    ["Megnézem, ki jön szombaton", "**Naptár** gomb, jobb felül (1. rész)"],
    ["Átírom egy levél szövegét", "**Beállítások** → a megfelelő alfül (1. rész)"],
    ["Válaszolok egy ajánlatkérésre", "**Ajánlatok** fül (3. rész)"],
    ["Kicserélem a főoldal képeit", "1. rész vége"],
    ["Lemondok egy egész alkalmat", "**Programok** → „Szerkesztés” → az időpont állapota: **Elmarad** (2. rész)"],
    ["Megnézem egy vendég számlázási címét", "**Foglalások** → „Szerkesztés”, illetve **Ajánlatok** → „Részletek” (2. és 3. rész)"],
    ["Megnézem, járt-e már nálunk valaki", "**Partnerek** gomb, jobb felül (1. rész)"],
    ["Valami nem működik", "1. rész: „Ha valami nem működik”"],
]),

# =====================================================================
("h1", "1. rész — Eligazodás és beállítások"),

("h2", "Belépés"),
("lepesek", [
    "Nyisd meg az oldal `/admin/` címét (pl. `kemence-akademia.xyz/admin/`).",
    "Írd be az e-mail címed és a jelszavad.",
    "Kilépés: jobb felül a **Kilépés** gomb.",
]),
("tipp", "**Regisztráció nincs**, és **minden admin fiók teljes jogú** — aki belép, mindent lát és mindent módosíthat. Új kolléga hozzáférését, vagy elfelejtett jelszó esetén az újat, a fejlesztő állítja be."),

("h2", "A felület felépítése"),
("p", "A fülsor a rendszer **két ágát** követi, mindkettőnek saját színe van:"),
("tabla", ["Csoport", "Fülek", "Mire való"], [
    ["**Nyilvános programok** (narancs)", "Foglalások, Programok", "Meghirdetett időpontok, amikre a vendég azonnal foglal."],
    ["**Egyedi megrendelések** (lila)", "Ajánlatok, Egyedi programok", "Időpont és ár nélküli egyedi programok (leánybúcsú, csapatépítő), amikre a vendég ajánlatot kér."],
    ["**Rendszer**", "⚙ Beállítások", "Csapat e-mail cím, azonosítók, levélsablonok."],
    ["jobb szélen", "👥 **Partnerek** · 📅 **Naptár**", "Mindkét ág egy helyen: a Naptár időrendben, a Partnerek pedig emberek szerint."],
]),
("tipp", "A logika mindkét ágban ugyanaz: a **Programok** és az **Egyedi programok** fülön azt állítod be, **mit kínálsz**; a **Foglalások** és az **Ajánlatok** fülön azt intézed, **ami beérkezett**."),
("kep", "az admin fülsora a két színes csoporttal"),

("h2", "A Naptár — ki jön és mikor"),
("lepesek", [
    "Kattints a jobb felső **Naptár** gombra.",
    "Lépj a kívánt hónapra vagy évre.",
    "Kattints egy alkalomra: látod a foglalókat névvel, létszámmal, elérhetőséggel.",
    "**📄 Jelentés (PDF)** — nyomtatható résztvevőlista.",
]),
("p", "A színek a telítettséget mutatják (**zöld** = van hely, **narancs** = majdnem tele, **piros** = betelt), a **lila** pedig az elfogadott egyedi ajánlat. Éves lista vagy havi rács nézet: a Beállításokban választod."),
("pelda", "Szombat reggel, indulás a helyszínre", [
    "**Naptár** → a szombati alkalomra kattintasz.",
    "Látod: „Nápolyi pizza este — 9/12 fő, 4 foglalás”.",
    "**📄 Jelentés (PDF)** → kinyomtatod: ki jön, hány fővel, milyen telefonszámon.",
    "A helyszínen a lista alapján fogadod a vendégeket.",
]),

("h2", "Partnerek — járt már nálunk ez az ember?"),
("p", "A **Partnerek** fülön egy listában szerepel mindenki, aki valaha **foglalt** egy programra vagy **ajánlatot kért**. Egy sor = egy ember: a rendszer az **e-mail cím** alapján vonja össze az alkalmait, így akkor is egy sorban látod, ha négyszer jelentkezett."),
("lista", [
    "**Alkalmak** oszlop: például „2 foglalás · 1 ajánlat”, mellette **Visszatérő** jelvény, ha egynél több alkalma volt.",
    "**Kereső minden oszlop fölött**: név, telefon, e-mail, lakcím és az alkalmak szerint.",
    "**Időszak**: szűkíthetsz egy adott időszakra (pl. csak az idei év).",
    "**Csak megvalósult alkalmak**: ha bekapcsolod, csak a jóváhagyott foglalások és az elfogadott ajánlatok számítanak.",
    "**✕ Szűrők törlése**: egy kattintással minden kereső és szűrő alaphelyzetbe áll. A gomb csak akkor jelenik meg, ha van mit törölni.",
    "**📊 Excel**: az éppen látott lista letölthető.",
]),
("tipp", "A sorra kattintva **lenyílnak az alkalmai**: azonosító, program vagy egyedi program, időpont, létszám, státusz. Az **azonosítóra kattintva** (pl. `F-101`) helyben megnyílik az adott foglalás szerkesztője, illetve az ajánlat részletei — nem kell fület váltanod, és bezárás után ugyanott folytatod, ahol abbahagytad."),
("pelda", "Telefonál valaki, hogy „tavaly már voltunk nálatok”", [
    "**Partnerek** gomb.",
    "A **Név** oszlop keresőjébe beírod: „Kovács”.",
    "Egy sor marad: Kovács Anna, „2 foglalás · 1 ajánlat”, **Visszatérő** jelvénnyel.",
    "Rákattintasz a sorra: látod, hogy tavaly júniusban lemondta, novemberben viszont ott volt, és decemberre ajánlatot is kért.",
]),
("figyelem", "A lista mindig a **meglévő** foglalásokból és ajánlatokból épül fel — nincs külön partner-nyilvántartás. Ha egy foglalást vagy ajánlatot véglegesen törölsz (🗑), az a partner is eltűnik innen. A törlő ablakok erre figyelmeztetnek is."),

("h2", "Beállítások"),
("p", "A **Beállítások** fül három alfülből áll:"),
("lista", [
    "**Általános** — csapat e-mail cím, a **két info-sáv** (a foglalási és az ajánlatkérő űrlap tetején megjelenő szöveg), és hogy a Naptár lista vagy rács nézetben nyíljon.",
    "**Nyilvános programok** — a foglalási azonosító formátuma (pl. `F-100`) és a foglalási levelek sablonjai.",
    "**Egyedi megrendelések** — az ajánlat-azonosító formátuma (pl. `A-100`) és az ajánlati levelek sablonjai.",
]),
("tipp", "A sablonokban a kapcsos zárójeles mezők (`{nev}`, `{program}`, `{idopont}`) küldéskor magától kitöltődnek. Minden kártyának **saját Mentés gombja** van — arra a kártyára vonatkozik, amelyikben van. Hogy melyik levél mikor megy ki, a **2. és a 3. részben** találod, a saját ágánál."),
("pelda", "Kiírom, hogy a számlát magánszemély nevére állítjuk", [
    "**Beállítások** → **Általános** alfül.",
    "Az **Ajánlatkérési info-sáv** mezőbe beírod a szöveget, pl. „A számlát magánszemély nevére állítjuk ki!”.",
    "Ugyanezt a **Foglalási info-sáv** mezőbe is beírhatod, ha a foglalási űrlapon is látszódjon.",
    "**Mentés** → a szöveg azonnal megjelenik az űrlap tetején. Ha kiüríted a mezőt, a sáv eltűnik.",
]),
("pelda", "Átírom a visszaigazoló levél szövegét", [
    "**Beállítások** → **Nyilvános programok** alfül.",
    "Görgess a „Foglalási e-mail sablonok” kártyához, keresd meg a „Visszaigazolás” blokkot.",
    "Írd át a törzsszöveget, de a `{nev}` és a `{program}` maradjon benne — ezeket a rendszer tölti ki.",
    "**Mentés**. A következő foglalásnál már az új szöveg megy ki.",
]),
("kep", "a Beállítások három alfüle"),

("h2", "Főoldali képek cseréje"),
("p", "A program-fotókat az adminban cseréled, a program szerkesztőjében. A főoldal fix képei (felső diavetítés, csapatkép, logó, workshop-képek) fájlok: a `web/kepek/` mappában vannak, az új képet **ugyanazon a néven** kell menteni, majd újra ki kell tenni az oldalt. Ideális méret: 1400–2000 képpont széles, 200–400 KB-os JPG."),

("h2", "Ha valami nem működik"),
("lista", [
    "**Nem látszik a módosításod** a főoldalon → frissíts **Ctrl+F5**-tel (a böngésző a régit tartotta meg).",
    "**Nem jött meg egy levél** → nézd meg a vendég sorában a „✉ N kiment levél” linket. Ha ott szerepel, a levél elment: a spam mappában lesz. Ha nem szerepel, szólj a fejlesztőnek.",
    "**Nektek nem jön értesítő** → Beállítások → Általános → a „Csapat e-mail cím” valódi, olvasott postafiók legyen.",
    "**A vendég nem tud foglalni** → az alkalom betelt, lejárt, elmaradt, vagy a foglalás szünetel (lásd 2. rész).",
    "**Elfelejtett jelszó** → a fejlesztő tud újat beállítani.",
]),
("tipp", "Az admin felület **telefonon is használható** — a listák keskeny képernyőn kártyákká rendeződnek. Hosszabb munkához (program felvitele, sablonszerkesztés) azért kényelmesebb a gép."),

# =====================================================================
("h1", "2. rész — Nyilvános programok és foglalások"),
("p", "Ez a rendszer nagyobbik fele: meghirdetsz egy alkalmat, a vendég foglal rá, te pedig jóváhagyod és levelet küldesz neki."),

("pelda", "Meghirdetek egy pizzás estét — az elejétől a végéig", [
    "**Programok** fül → **+ Új program**.",
    "Cím: „Nápolyi pizza este”. Rövid leírás: „Tésztától a kemencéig — egy estébe sűrítve.”",
    "Állapot: **Aktív**.",
    "**+ Időpont hozzáadása** → dátum: „2026-11-08 18:00”, ár: „15 000”, max. létszám: „12”.",
    "**Mentés** → a program azonnal megjelenik a főoldalon, foglalható időponttal.",
    "Pár nap múlva érkezik az első foglalás, és a **Foglalások** fülön már intézheted is.",
]),

("h2", "Új programot hirdetnék meg"),
("lepesek", [
    "**Programok** fül → **+ Új program**.",
    "Cím és rövid leírás (ez látszik a kártyán). Ha van: részletes leírás, előadó, várható időtartam, fotó.",
    "Állapot: **Aktív**, ha foglalható. **Hamarosan**, ha csak beharangoznád (ilyenkor nincs időpont).",
    "**+ Időpont hozzáadása** — alkalmanként: dátum és idő, ár/fő, kedvezményes ár (nem kötelező), max. létszám.",
    "**Mentés**.",
]),
("tipp", "Egy program **egy kártya** a főoldalon, akárhány időponttal. Ha ugyanazt a workshopot többször tartod, ne vegyél fel új programot — adj hozzá időpontot."),

("h2", "Új időpontot adnék hozzá, vagy módosítanék"),
("lepesek", [
    "A programnál **Szerkesztés**.",
    "Új alkalomhoz: **+ Időpont hozzáadása**. Módosításhoz írd át a mezőket (dátum, ár, létszám).",
    "**Mentés**.",
]),
("tipp", "A max. létszámot nem viheted a már lefoglalt létszám alá — a rendszer szól, és visszaállítja az eredeti értéket."),

("h2", "Zártkörű alkalmat hirdetnék (csoportoknak)"),
("p", "Az időpont sorában két nem kötelező mező van:"),
("lista", [
    "**Min. fő/foglalás** — egy foglalás legalább ennyi fős lehet.",
    "**Max. foglalás** — összesen ennyi foglalás (csapat) jöhet, akármekkorák.",
]),
("pelda", "Céges csapatépítő: az egész műhely egy csapaté", [
    "A programnál **Szerkesztés** → a november 20-i időpont sorában:",
    "**Max. foglalás** = `1` → csak egyetlen csapat jöhet.",
    "**Min. fő/foglalás** = `10` → de legalább 10 fősnek kell lennie.",
    "**Mentés**. Ezután egy 10 fős csapat lefoglalja az egész estét, és az alkalom lezár.",
    "Ha csak a „Max. foglalás = 1” lenne beállítva, egy **2 fős** foglalás is lezárná a teljes alkalmat.",
]),
("tipp", "A két mező **független**, bármelyik megadható a másik nélkül. Részletes magyarázat és további példák: az időpont-mezők alatti **?** gomb."),

("h2", "Elmarad egy alkalom — lemondás és értesítés mindenkinek"),
("p", "Nincs külön „lemondom az egészet” gomb: az **időpont állapotán** keresztül megy, és a rendszer intézi a foglalásokat is."),
("lepesek", [
    "A programnál **Szerkesztés** → az adott időpont állapota: **Elmarad** → **Mentés**.",
    "Ha van rajta élő foglalás, a rendszer megáll és megkérdezi, hogyan folytassa.",
    "Három választásod van: **Mégse** (mégsem marad el) · **Lemondás, e-mail értesítő nélkül** · **Lemondás + elmaradás e-mail értesítő** — ez utóbbi egy lépésben, mindenkinek kiküldi a „sajnos elmarad” levelet.",
]),
("tipp", "Az elmaradás-levél szövegét előre a **Beállítások → Nyilvános programok → Foglalási e-mail sablonok** alatt szerkeszted, mert tömegesen megy ki — itt nincs egyenkénti szerkesztés, mint a **✉ Levél** gombnál."),
("tipp", "Foglalással rendelkező időpontot **ne törölj** — az „Elmarad” megőrzi a nyomát és értesíti a vendégeket. Törölni (✕) csak foglalás nélküli időpontot érdemes."),

("h2", "Lejártak az időpontok — mi látszik kint?"),
("p", "A főoldal csak a **jövőbeli** alkalmakat mutatja. Ha egy programnak nincs ilyenje, a kártyája ott marad, de foglalni nem lehet rajta:"),
("lista", [
    "**Hamarosan** jelvény — ha a program állapota „Hamarosan” (beharangozás, még nincs dátum);",
    "**Új időpontok egyeztetés alatt** sáv — ha a program „Aktív”, de nincs jövőbeli időpontja.",
]),
("tipp", "Ha tervezed újra megtartani: **Szerkesztés → + Időpont hozzáadása**. Ha nem: **archiváld**, különben ott marad a főoldalon."),

("h2", "Egy időre leállítanám a foglalást"),
("lista", [
    "**Egy programra:** a programnál **Foglalás felfüggesztése**. A kártya látszik, de nem foglalható. Vissza: **Foglalás engedélyezése**.",
    "**Mindenre:** a Programok fül tetején **Minden foglalás felfüggesztése**.",
]),
("tipp", "Akkor hasznos, ha még nem biztos, hogy elindul egy program, vagy épp átírod az adatait. A meglévő foglalások megmaradnak."),

("h2", "Lement a program — archiválás vagy törlés"),
("lista", [
    "**Archiválás** — lekerül a főoldalról, de az adatok és a foglalások megmaradnak („Archivált programok” alfül). Bármikor visszaállítható. **Általában ezt válaszd.**",
    "**Törlés** — végleg eltűnik minden időpontjával. Csak foglalás nélküli programnál jelenik meg.",
]),
("p", "A főoldali sorrendet a programlistában a **⠿** fogantyúnál húzva állítod be."),

("h2", "Egy foglalás útja"),
("tabla", ["Státusz", "Mit jelent", "Foglal helyet?"], [
    ["**Jóváhagyásra vár**", "Új foglalás, még nem döntöttél.", "Igen"],
    ["**Jóváhagyott**", "Elfogadtad, a hely biztos.", "Igen"],
    ["**Elutasított**", "Nem fogadtad el.", "Nem — a hely felszabadul"],
    ["**Lemondott**", "A vendég lemondta.", "Nem — a hely felszabadul"],
]),
("p", "A **Foglalások** fülön minden státusz külön alfül (zárójelben a darabszám); az „Összes foglalás” mindent mutat. Szűrhetsz időpontra, névre, és arra, hogy jövőbeli / mai / múltbéli."),

("h2", "Új foglalás érkezett — mit csináljak?"),
("lepesek", [
    "**Foglalások** → **Jóváhagyásra vár** alfül.",
    "Nézd meg a foglalást (a 📝 jel = a vendég írt megjegyzést).",
    "Ha rendben van: **Jóváhagyás**.",
    "Utána a sorban **✉ Levél** → a jóváhagyó levél. Küldés előtt átírhatod a szöveget.",
    "Ha nem tudod fogadni: **Elutasítás**, majd **✉ Levél** az elutasítóval.",
]),
("pelda", "Kovács Anna foglalása, az elejétől a végéig", [
    "Anna 4 főre foglal a november 8-i pizzás estére. **Azonnal kap** egy automatikus visszaigazolást, ti pedig értesítőt a csapat címére.",
    "A foglalás „Jóváhagyásra vár” állapotban áll. Megérkezik az utalás.",
    "**Jóváhagyás** → a státusz „Jóváhagyott” lesz.",
    "**✉ Levél** → elküldöd a jóváhagyó levelet; beleírod, hogy hozzon kötényt.",
    "November 7-én a rendszer **magától** küld neki emlékeztetőt.",
    "November 8-án reggel: **Naptár** → az alkalom → **📄 Jelentés (PDF)**, és Anna neve rajta van a listán.",
]),
("tipp", "**Miért két lépés a jóváhagyás és a levél?** Mert a státusz átállítása magától nem küld levelet. Így előbb döntesz, és csak utána, személyre szabott szöveggel értesíted a vendéget. Ha tévedésből hagytál jóvá: **Jóváhagyás visszavonása**."),
("tipp", "**Számlázási cím:** a vendég a foglaláskor megadja az **irányítószámot**, a **helységet** és a **további címadatot** — enélkül nem tudja elküldeni az űrlapot. A listában nem látszik (hogy ne legyen zsúfolt), de a **Szerkesztés** ablakban megnézheted és javíthatod is."),
("kep", "a Foglalások lista egy sorral és a gombokkal"),

("h2", "Lemondás, módosítás, áthelyezés"),
("lista", [
    "**A vendég lemondta** → a foglalásnál „Vendég lemondta”, majd ha kell, **✉ Levél**. A hely azonnal felszabadul.",
    "**Más létszám vagy adat** → „Szerkesztés”, átírod, **Mentés**.",
    "**Másik időpontra menne** → „Szerkesztés”, és ugyanannak a programnak egy másik időpontját választod.",
]),
("tipp", "Áthelyezni csak **azonos árú** időpontra lehet — eltérő árnál a vendégnek újra kell foglalnia. A férőhelyet a rendszer itt is betartja."),

("h2", "Milyen levelek mennek ki a foglalásoknál?"),
("tabla", ["Levél", "Mikor", "Hogyan"], [
    ["Visszaigazolás a vendégnek („megkaptuk”)", "amint leadja a foglalást", "automatikus"],
    ["Értesítő nektek, a csapat címére", "amint leadja a foglalást", "automatikus"],
    ["Jóváhagyás / elutasítás / lemondás", "miután átállítod a státuszt", "te küldöd: **✉ Levél**"],
    ["Emlékeztető a vendégnek", "a program előtti napon", "automatikus"],
    ["„Sajnos elmarad”", "amikor egy időpontot elmaradóra állítasz", "választható, mindenkinek egyszerre"],
]),
("tipp", "A vendég **válasza a csapat e-mail címre** érkezik. A foglalás sorában a „**✉ N kiment levél**” linkre kattintva látod, mi ment már ki neki — ez a leggyorsabb ellenőrzés, ha a vendég azt mondja, nem kapott semmit. A levelek szövegét a **Beállítások → Nyilvános programok** alfülön írhatod át."),

("h2", "Résztvevőlista, Excel, takarítás"),
("lista", [
    "**📄 Jelentés (PDF)** — a Naptárban, egy alkalomra kattintva: nyomtatható résztvevőlista.",
    "**📊 Excel** — a Foglalások fülön az éppen látott listát tölti le (az aktív alfül és a szűrők szerint). A **számlázási cím** külön oszlopokban benne van, így számlázáshoz is jó.",
    "**🗑 Törlés** — csak elutasított vagy lemondott foglalásnál jelenik meg. Élő foglalást nem lehet törölni.",
]),

# =====================================================================
("h1", "3. rész — Egyedi megrendelések (ajánlatok)"),
("p", "Az egyedi ág **időpont és ár nélküli egyedi programokat** kínál (leánybúcsú, céges csapatépítő, gyerekzsúr). A vendég ezekre **ajánlatot kér**, az árat és az időpontot pedig e-mailben egyeztetitek. A rendszer azt tartja nyilván, hol tart az ügy."),

("pelda", "Leánybúcsú ajánlatkérés — az elejétől a végéig", [
    "Szabó Rita ajánlatot kér a „Leánybúcsú a kemencénél” egyedi programra: 14 fő, november 22., „szombat délután szeretnénk, pezsgővel”.",
    "Rita **azonnal kap** visszaigazolást, ti értesítőt. Az ügy „Ajánlatra vár” állapotba kerül.",
    "**Ajánlatok** fül → **Részletek**: elolvasod a kérést. A **belső jegyzetbe** felírod: „„kérdezd meg a tortát”” — ezt a vendég soha nem látja.",
    "Összeállítod az ajánlatot, és **a saját leveleződből** elküldöd neki: 14 fő, 210 000 Ft, 15:00-tól.",
    "Visszajössz az adminba: **Ajánlat kiküldve** — így látszik, hogy a válaszára vársz.",
    "Rita igent mond. **Elfogad** → végleges időpont: „2026-11-22 15:00”, létszám: „14”, összár: „210 000”, az ajánlat szövege bemásolva.",
    "**✉ Levél** → megerősítő levél a végleges adatokkal.",
    "A rendezvény ettől kezdve **lilával látszik a Naptárban**, a foglalások mellett.",
]),

("h2", "Új egyedi programot teszek ki a főoldalra"),
("lepesek", [
    "**Egyedi programok** fül → **+ Új egyedi program**.",
    "Cím, rövid leírás (a kártyára), ha kell részletes leírás és fotó.",
    "**Mentés**. A kártyán a vendég a **„Kérjen egyedi ajánlatot”** gombot látja.",
]),
("tipp", "Nincs rajta időpont és ár — ez a lényege. Ha már nem kínálod: **Archiválás**; a beérkezett kérések megmaradnak. Sorrend: a **⠿** fogantyúnál húzd."),
("tipp", "**Törölni csak olyan egyedi programot lehet, amire még nem érkezett ajánlatkérés** — ugyanaz a szabály, mint a programoknál. Ha már jött rá kérés, a kártyán ott a darabszám, és csak az **Archiválás** marad. Így az Ajánlatok listájában sosem vész el, hogy a vendég mire kért ajánlatot."),

("h2", "Ajánlatkérés érkezett — a menet"),
("lepesek", [
    "**Ajánlatok** fül → **Részletek**: létszám, kívánt időpont, üzenet. Ide írhatsz **belső jegyzetet** is.",
    "Állítsd össze az ajánlatot, és küldd el **a saját leveleződből**.",
    "Jelöld a rendszerben: **Ajánlat kiküldve**.",
    "Ha elfogadta: **Elfogad** → végleges időpont, létszám, **összár** (a teljes rendezvényre, nem fejenként) és az ajánlat szövege. Fájlt is csatolhatsz.",
    "**✉ Levél** → megerősítés a végleges adatokkal (a csatolmány is megy vele).",
]),
("tipp", "**Miért érdemes az adminban is elfogadni?** Mert az elfogadott ajánlat bekerül a **Naptárba**, így egy helyen látod a foglalásokkal — és nem foglalsz rá véletlenül másik rendezvényt."),
("tipp", "**Számlázási cím:** az ajánlatkérő is megadja (irányítószám, helység, további címadat). A **Részletek** ablakban látod, és ott javíthatod is — a belső jegyzettel együtt, egy **Mentés** gombbal."),
("kep", "az Ajánlatok lista és az Elfogad ablak"),

("h2", "Változott az időpont vagy az ár"),
("lepesek", [
    "Az ajánlatnál **Részletek** → **Végleges adatok módosítása / megtekintése**.",
    "Írd át, **Mentés**.",
    "Ha kell, küldj új megerősítést: **✉ Levél**.",
]),

("h2", "Nem lesz belőle rendezvény"),
("lista", [
    "**Elutasítás** — ha te nem vállalod.",
    "**Lemondás** — ha a vendég lép vissza (elfogadott ajánlatnál is).",
    "Mindkettő után küldhetsz levelet (**✉ Levél**), a lezárt ajánlatokat pedig a **🗑 Törlés** takarítja.",
]),

("h2", "Milyen levelek mennek ki az ajánlatoknál?"),
("tabla", ["Levél", "Mikor", "Hogyan"], [
    ["Visszaigazolás a vendégnek („megkaptuk a kérésed”)", "amint beküldi az ajánlatkérést", "automatikus"],
    ["Értesítő nektek, a csapat címére", "amint beküldi az ajánlatkérést", "automatikus"],
    ["**Maga az ajánlat**", "amikor összeállítottad", "a **saját leveleződből**, a rendszeren kívül"],
    ["Megerősítés a végleges adatokkal", "miután **Elfogad**-tad", "te küldöd: **✉ Levél** (a csatolmánnyal)"],
    ["Elutasítás / lemondás", "miután átállítod a státuszt", "te küldöd: **✉ Levél**"],
]),
("tipp", "Az **ajánlatot magát a rendszer nem küldi el** — azt te írod meg a saját leveleződben, mert minden ajánlat más. A rendszer azt tartja nyilván, hogy hol tart az ügy, és a végleges adatokról küld megerősítést. A szövegeket a **Beállítások → Egyedi megrendelések** alfülön írhatod át."),

("h2", "Hol látom, mi van folyamatban?"),
("p", "Az **Ajánlatok** fülön a státuszok külön alfülek, zárójelben a darabszámmal: „Ajánlatra vár”, „Kiküldve”, „Elfogadva”, „Elutasítva”, „Lemondva”, „Összes”. Az elfogadott rendezvények a **Naptárban** is megjelennek lilával, és a **📊 Excel** gombbal az éppen látott lista letölthető — a számlázási címmel együtt."),

("osszegzes", [
    "**Foglalás:** értesítő e-mail → **Jóváhagyás** + **✉ Levél** → előző nap automatikus emlékeztető → Naptár → **📄 Jelentés (PDF)** a helyszínre.",
    "**Ajánlat:** értesítő e-mail → **Részletek** → ajánlat a saját leveleződből → **Ajánlat kiküldve** → **Elfogad** → **✉ Levél** → lilával a Naptárban.",
]),
]


# =====================================================================
#  Inline jelölés:  **félkövér**,  `kód`
# =====================================================================
JELOLES = re.compile(r"(\*\*.+?\*\*|`.+?`)")

def darabol(szoveg):
    """[(szöveg, 'b'|'code'|None), ...]"""
    ki = []
    for d in JELOLES.split(szoveg):
        if not d: continue
        if d.startswith("**") and d.endswith("**"): ki.append((d[2:-2], "b"))
        elif d.startswith("`") and d.endswith("`"): ki.append((d[1:-1], "code"))
        else: ki.append((d, None))
    return ki

def inline_html(szoveg):
    out = []
    for t, s in darabol(szoveg):
        e = html.escape(t, quote=False)
        out.append(f"<b>{e}</b>" if s == "b" else f"<code>{e}</code>" if s == "code" else e)
    return "".join(out)


# =====================================================================
#  HTML
# =====================================================================
CSS = """
  :root{--bg:#faf6f0; --card:#fff; --ink:#241a12; --muted:#7a6a58; --line:#e7ddd0;
    --accent:#e4571b; --accent-2:#b8430f; --egyedi:#8a5cc7; --pelda:#2f7d4f; --code:#f3ece2}
  *{box-sizing:border-box}
  body{margin:0; background:var(--bg); color:var(--ink);
    font-family:"Segoe UI",system-ui,-apple-system,sans-serif; line-height:1.6; font-size:16px}
  .wrap{max-width:860px; margin:0 auto; padding:32px 22px 80px}
  header h1{font-size:30px; margin:0 0 6px}
  header p{color:var(--muted); margin:0 0 4px}
  .frissitve{font-size:13px}
  h1.fejezet{font-size:24px; margin:44px 0 10px; padding-bottom:6px; border-bottom:2px solid var(--line)}
  h2{font-size:18px; margin:26px 0 6px; color:var(--accent-2)}
  p{margin:8px 0}
  code{background:var(--code); padding:1px 6px; border-radius:5px; font-family:"Cascadia Code",Consolas,monospace; font-size:14px}
  ol,ul{margin:8px 0; padding-left:24px} li{margin:4px 0}
  table{border-collapse:collapse; width:100%; margin:12px 0; background:var(--card)}
  th,td{border:1px solid var(--line); padding:7px 11px; text-align:left; vertical-align:top}
  th{background:#f6efe6; font-size:14px}
  .doboz{border-left:4px solid var(--accent); background:#fff4ee; padding:10px 15px; border-radius:8px; margin:12px 0}
  .doboz.tipp{border-color:#d9a441; background:#fdf7e3}
  .doboz.osszeg{border-color:var(--egyedi); background:#f5effb}
  .doboz.pelda{border-color:var(--pelda); background:#eef6ef}
  .doboz.pelda .pelda-cim{display:block; margin-bottom:4px; color:#25623e}
  .doboz ul,.doboz ol{margin:4px 0}
  .kep{border:2px dashed #c9b79f; background:#faf3e9; color:#8a7458; border-radius:10px;
    padding:16px; margin:12px 0; text-align:center; font-size:14px}
  .tartalom{margin:18px 0}
  .tartalom a{display:block; padding:2px 0; text-decoration:none; color:var(--accent-2)}
  .tartalom .toc-resz{font-weight:700; margin-top:10px}
  .tartalom .toc-helyzet{padding-left:22px; font-size:15px; color:var(--muted)}
  .tartalom .toc-helyzet:hover{color:var(--accent-2)}
  @media print{ body{background:#fff} .kep{display:none} }
"""

def html_gyart():
    t = []
    t.append(f'<header><h1>🔥 {html.escape(CIM)}</h1><p>{html.escape(ALCIM)}</p>'
             f'<p class="frissitve">{FRISSITVE} · A 📷 jelölés a később beillesztendő képernyőképek helye.</p></header>')
    for i, b in enumerate(TARTALOM):
        tip = b[0]
        if tip == "bevezeto":
            t.append('<div class="doboz"><b>A rendszer két része</b><ul>'
                     + "".join(f"<li>{inline_html(x)}</li>" for x in b[1]) + "</ul></div>")
            sorok = []
            for j, x in enumerate(TARTALOM):
                if x[0] == "h1":
                    sorok.append(f'<a class="toc-resz" href="#f{j}">{html.escape(x[1])}</a>')
                elif x[0] == "h2":
                    sorok.append(f'<a class="toc-helyzet" href="#f{j}">{html.escape(x[1])}</a>')
            t.append('<div class="tartalom"><b>Tartalom</b>' + "".join(sorok) + "</div>")
        elif tip == "h1":        t.append(f'<h1 class="fejezet" id="f{i}">{html.escape(b[1])}</h1>')
        elif tip == "h2":        t.append(f'<h2 id="f{i}">{html.escape(b[1])}</h2>')
        elif tip == "p":         t.append(f"<p>{inline_html(b[1])}</p>")
        elif tip == "lepesek":   t.append("<ol>" + "".join(f"<li>{inline_html(x)}</li>" for x in b[1]) + "</ol>")
        elif tip == "lista":     t.append("<ul>" + "".join(f"<li>{inline_html(x)}</li>" for x in b[1]) + "</ul>")
        elif tip == "tipp":      t.append(f'<div class="doboz tipp">💡 {inline_html(b[1])}</div>')
        elif tip == "figyelem":  t.append(f'<div class="doboz">⚠️ {inline_html(b[1])}</div>')
        elif tip == "kep":       t.append(f'<div class="kep">📷 <b>Képernyőkép:</b> {html.escape(b[1])}</div>')
        elif tip == "pelda":
            t.append('<div class="doboz pelda"><b class="pelda-cim">📋 Példa — ' + html.escape(b[1]) + "</b><ol>"
                     + "".join(f"<li>{inline_html(x)}</li>" for x in b[2]) + "</ol></div>")
        elif tip == "tabla":
            fej, sorok = b[1], b[2]
            t.append("<table><tr>" + "".join(f"<th>{html.escape(h)}</th>" for h in fej) + "</tr>"
                     + "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in s) + "</tr>" for s in sorok)
                     + "</table>")
        elif tip == "osszegzes":
            t.append('<div class="doboz osszeg"><b>Röviden, a napi rutin</b><ul>'
                     + "".join(f"<li>{inline_html(x)}</li>" for x in b[1]) + "</ul></div>")
    return ("<!doctype html>\n<html lang=\"hu\">\n<head>\n<meta charset=\"utf-8\">\n"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
            f"<title>{html.escape(CIM)}</title>\n"
            "<!-- GENERÁLT FÁJL — ne kézzel szerkeszd! Forrás: build_docx.py (TARTALOM) -->\n"
            f"<style>{CSS}</style>\n</head>\n<body>\n<div class=\"wrap\">\n"
            + "\n".join(t) + "\n</div>\n</body>\n</html>\n")


# =====================================================================
#  DOCX
# =====================================================================
NARANCS = RGBColor(0xB8, 0x43, 0x0F)
SZURKE  = RGBColor(0x7A, 0x6A, 0x58)

def futasok(par, szoveg, alap_bold=False):
    for t, s in darabol(szoveg):
        r = par.add_run(t)
        r.bold = alap_bold or s == "b"
        if s == "code":
            r.font.name = "Consolas"; r.font.size = Pt(10)

def arnyekol(par, szin, keret):
    """Bekezdés háttérszín + bal oldali színes csík (a HTML-es dobozok megfelelője)."""
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bal = OxmlElement("w:left")
    bal.set(qn("w:val"), "single"); bal.set(qn("w:sz"), "24"); bal.set(qn("w:space"), "8"); bal.set(qn("w:color"), keret)
    bdr.append(bal); pPr.append(bdr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), szin)
    pPr.append(shd)

def cella_hatter(cella, szin):
    tcPr = cella._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), szin)
    tcPr.append(shd)

def doboz(doc, cim, elemek, szin, keret):
    if cim:
        p = doc.add_paragraph(); futasok(p, cim, alap_bold=True); arnyekol(p, szin, keret)
        p.paragraph_format.space_after = Pt(0)
    for x in elemek:
        p = doc.add_paragraph(); futasok(p, x); arnyekol(p, szin, keret)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def toc_mezo(doc, szoveg_ha_nincs_frissitve):
    """Valódi Word tartalomjegyzék-MEZŐ (nem statikus lista).

    A Heading 1/2 stílusú címekből épül; Ctrl+kattintásra a fejezetre ugrik,
    és oldalszámokat is mutat. A Word a megnyitáskor frissíti (lásd
    `frissitsd_a_mezoket`), vagy kézzel: kattints rá és nyomj F9-et.
    """
    p = doc.add_paragraph()
    def elemet_ad(elem):
        run = p.add_run()
        run._r.append(elem)

    kezd = OxmlElement("w:fldChar")
    kezd.set(qn("w:fldCharType"), "begin"); kezd.set(qn("w:dirty"), "true")
    elemet_ad(kezd)

    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-2" \h \z \u '
    elemet_ad(instr)

    elvalaszt = OxmlElement("w:fldChar"); elvalaszt.set(qn("w:fldCharType"), "separate")
    elemet_ad(elvalaszt)

    # Ez látszik addig, amíg a Word nem frissítette a mezőt:
    jelolo = p.add_run(szoveg_ha_nincs_frissitve)
    jelolo.italic = True; jelolo.font.color.rgb = SZURKE

    veg = OxmlElement("w:fldChar"); veg.set(qn("w:fldCharType"), "end")
    elemet_ad(veg)
    return p


def frissitsd_a_mezoket(doc):
    """A Word a megnyitáskor magától frissítse a tartalomjegyzéket."""
    try:
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        doc.settings.element.append(uf)
    except Exception:
        pass   # ha nem sikerül, a kézi frissítés (F9) akkor is működik


def docx_gyart():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Segoe UI"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
    for nev, meret in (("Heading 1", 16), ("Heading 2", 12.5)):
        s = doc.styles[nev]; s.font.name = "Segoe UI"; s.font.size = Pt(meret); s.font.color.rgb = NARANCS
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)

    c = doc.add_heading(CIM, level=0)
    for r in c.runs: r.font.size = Pt(22); r.font.color.rgb = NARANCS
    p = doc.add_paragraph(ALCIM); p.runs[0].font.color.rgb = SZURKE
    p = doc.add_paragraph(FRISSITVE + " · A 📷 jelölés a később beillesztendő képernyőképek helye.")
    p.runs[0].font.size = Pt(9); p.runs[0].font.color.rgb = SZURKE

    for b in TARTALOM:
        tip = b[0]
        if tip == "bevezeto":
            doboz(doc, "A rendszer két része", ["• " + x for x in b[1]], "FFF4EE", "E4571B")
            p = doc.add_paragraph(); p.add_run("Tartalom").bold = True
            toc_mezo(doc, "(A tartalomjegyzék frissítéséhez: kattints ide, majd nyomj F9-et.)")
        elif tip == "h1":  doc.add_heading(b[1], level=1)
        elif tip == "h2":  doc.add_heading(b[1], level=2)
        elif tip == "p":
            p = doc.add_paragraph(); futasok(p, b[1])
        elif tip in ("lepesek", "lista"):
            # Kézi számozás: a Word „List Number” stílusa a listák között nem kezdi újra a számozást.
            for n, x in enumerate(b[1], 1):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Cm(0.9); pf.first_line_indent = Cm(-0.6); pf.space_after = Pt(2)
                p.add_run(f"{n}.\t" if tip == "lepesek" else "•\t")
                futasok(p, x)
        elif tip == "tipp":      doboz(doc, None, ["💡 " + b[1]], "FDF7E3", "D9A441")
        elif tip == "figyelem":  doboz(doc, None, ["⚠️ " + b[1]], "FFF4EE", "E4571B")
        elif tip == "pelda":
            doboz(doc, "📋 Példa — " + b[1],
                  [f"{n}. {x}" for n, x in enumerate(b[2], 1)], "EEF6EF", "2F7D4F")
        elif tip == "kep":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("📷 Képernyőkép: " + b[1]); r.italic = True; r.font.color.rgb = SZURKE
        elif tip == "tabla":
            fej, sorok = b[1], b[2]
            t = doc.add_table(rows=1, cols=len(fej)); t.style = "Table Grid"
            for i, h in enumerate(fej):
                cl = t.rows[0].cells[i]; cl.text = ""
                futasok(cl.paragraphs[0], h, alap_bold=True); cella_hatter(cl, "F6EFE6")
            for s in sorok:
                cells = t.add_row().cells
                for i, x in enumerate(s):
                    cells[i].text = ""; futasok(cells[i].paragraphs[0], x)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif tip == "osszegzes":
            doboz(doc, "Röviden, a napi rutin", ["• " + x for x in b[1]], "F5EFFB", "8A5CC7")
    frissitsd_a_mezoket(doc)
    return doc


if __name__ == "__main__":
    with open(os.path.join(GYOKER, "felhasznaloi-utmutato.html"), "w", encoding="utf-8", newline="\n") as f:
        f.write(html_gyart())
    docx_gyart().save(os.path.join(GYOKER, "felhasznaloi-utmutato.docx"))
    print("Kész: felhasznaloi-utmutato.html + felhasznaloi-utmutato.docx")
